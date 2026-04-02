"""
scan_word.py — Word (.docx) Accessibility Scanner
===================================================
Persistent utility for Word document accessibility audits.
Checks against Microsoft Accessibility Checker rules + WCAG 2.2 AA.

Checks:
  Metadata:
    - Document title (core_properties.title)          [DOCX-META.TITLE]
    - Document language (core_properties.language)    [DOCX-META.LANG]

  Structure:
    - Heading structure (Heading 1/2/3 style usage)   [DOCX-STRUCT.HEADINGS]
    - Skipped heading levels                           [DOCX-STRUCT.HEADINGSKIP]
    - Empty paragraphs used as spacing                [DOCX-STRUCT.EMPTYPARA]
    - Table of Contents in long documents             [DOCX-STRUCT.TOC]

  Images & media:
    - Inline images missing alt text (descr attr)    [DOCX-IMG.ALT]
    - Floating images/shapes missing alt text         [DOCX-IMG.ALT_FLOAT]
    - Image alt text quality (placeholder text)       [DOCX-IMG.ALT_QUALITY]

  Tables:
    - Tables missing header row style                 [DOCX-TABLE.HEADERS]
    - Merged cells detection                          [DOCX-TABLE.MERGE]
    - Nested tables detection                         [DOCX-TABLE.NESTED]

  Lists:
    - Manual bullets vs semantic list styles          [DOCX-LIST.SEMANTIC]

  Links:
    - Hyperlinks with non-descriptive text            [DOCX-LINK.DESCRIPTIVE]

  Document health:
    - Tracked changes left in document                [DOCX-REVIEW.TRACKED]
    - Header/footer meaningful content                [DOCX-STRUCT.HDRFTR]
    - Footnotes / endnotes accessibility              [DOCX-STRUCT.FOOTNOTES]

Usage:
    python tools/scan_word.py <docx_or_folder> [--json] [--output file.json]

Dependencies: python-docx (pip install python-docx), lxml

This file is PERMANENT — do not delete. Used by agents and scan_all.py.
"""

import sys
import json
import argparse
import re
from pathlib import Path
from collections import Counter
from lxml import etree
from zipfile import ZipFile, BadZipFile

try:
    from docx import Document
    from docx.oxml.ns import qn
    import docx
except ImportError:
    sys.exit("ERROR: python-docx not installed. Run: pip install python-docx")


# ---------------------------------------------------------------------------
# Namespace helpers
# ---------------------------------------------------------------------------

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
WPS_NS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
WPC_NS = "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas"
WPG_NS = "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"
PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"

VAGUE_LINK_PATTERNS = re.compile(
    r"^(click here|here|read more|learn more|more|link|more info|"
    r"this|download|go|view|see here|details|info)\.?$",
    re.IGNORECASE,
)

HEADING_STYLES = {f"Heading {i}" for i in range(1, 10)}


# ---------------------------------------------------------------------------
# Analysis functions
# ---------------------------------------------------------------------------

def check_metadata(doc) -> list:
    findings = []
    cp = doc.core_properties

    title = cp.title if cp.title else ""
    if not title.strip():
        findings.append({
            "rule": "DOCX-META.TITLE",
            "severity": "Error",
            "confidence": "High",
            "wcag": "2.4.2",
            "message": "Document title is not set (File > Properties > Title).",
            "fix": "Set the document title via File > Info > Properties > Title in Word.",
        })

    language = cp.language if cp.language else ""
    if not language.strip():
        findings.append({
            "rule": "DOCX-META.LANG",
            "severity": "Error",
            "confidence": "High",
            "wcag": "3.1.1",
            "message": "Document language is not set.",
            "fix": "In Word: Review → Language → Set Proofing Language → set language for all text.",
        })

    return findings


def check_headings(doc) -> list:
    findings = []
    seen_levels = []

    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        for i in range(1, 7):
            if style_name == f"Heading {i}":
                seen_levels.append(i)
                break

    if not seen_levels:
        # Check if there's any meaningful body text to determine if headings are expected
        body_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        if len(body_texts) > 10:
            findings.append({
                "rule": "DOCX-STRUCT.HEADINGS",
                "severity": "Warning",
                "confidence": "Medium",
                "wcag": "1.3.1",
                "message": "No heading styles used. Long documents require structural headings for navigation.",
                "fix": "Apply Heading 1/2/3 styles from Word's Styles panel instead of bold/large text.",
            })
        return findings

    # Check that the first heading is H1
    if seen_levels[0] != 1:
        findings.append({
            "rule": "DOCX-STRUCT.HEADINGSKIP",
            "severity": "Warning",
            "confidence": "High",
            "wcag": "1.3.1",
            "message": (
                f"First heading in document is H{seen_levels[0]} instead of H1. "
                "Screen readers expect the document to start with H1."
            ),
            "fix": "Ensure the first heading in the document uses the Heading 1 style.",
        })

    # Check for skipped levels (track all levels seen so far)
    seen_set = set()
    for level in seen_levels:
        for expected in range(1, level):
            if expected not in seen_set:
                findings.append({
                    "rule": "DOCX-STRUCT.HEADINGSKIP",
                    "severity": "Warning",
                    "confidence": "High",
                    "wcag": "1.3.1",
                    "message": (
                        f"Heading level skipped: H{level} used before any H{expected} appears. "
                        "Screen readers navigate headings sequentially."
                    ),
                    "fix": "Ensure heading levels are consecutive: H1 → H2 → H3 without gaps.",
                })
                break  # report only the first missing level per heading
        seen_set.add(level)

    return findings


def check_images(doc) -> list:
    findings = []
    missing_alt = []
    empty_alt = []

    for shape in doc.inline_shapes:
        # python-docx inline shape
        # Alt text lives in the drawing's docPr element: wp:docPr @descr on the inline element
        try:
            inline_el = shape._inline
            # Look for wp:docPr with descr or title
            doc_pr_list = inline_el.findall(
                f".//{{{WP_NS}}}docPr"
            )
            if not doc_pr_list:
                missing_alt.append("(shape with no docPr)")
                continue
            for doc_pr in doc_pr_list:
                descr = doc_pr.get("descr", "")
                title = doc_pr.get("title", "")
                name = doc_pr.get("name", "image")
                if not descr.strip() and not title.strip():
                    missing_alt.append(name)
                elif descr.strip() and len(descr.strip()) < 4:
                    empty_alt.append(f"{name}: '{descr}'")
        except Exception:
            missing_alt.append("(shape inspection error)")

    if missing_alt:
        findings.append({
            "rule": "DOCX-IMG.ALT",
            "severity": "Error",
            "confidence": "High",
            "wcag": "1.1.1",
            "message": (
                f"{len(missing_alt)} image(s) have no alt text: "
                f"{missing_alt[:5]}{'...' if len(missing_alt) > 5 else ''}"
            ),
            "fix": (
                "Right-click each image in Word → View Alt Text → "
                "enter a description of the image's content and purpose. "
                "If decorative, check 'Mark as decorative'."
            ),
        })

    if empty_alt:
        findings.append({
            "rule": "DOCX-IMG.ALT_QUALITY",
            "severity": "Warning",
            "confidence": "Medium",
            "wcag": "1.1.1",
            "message": f"{len(empty_alt)} image(s) have very short alt text (likely placeholder): {empty_alt}",
            "fix": "Review alt text to ensure it meaningfully describes the image content.",
        })

    return findings


def check_floating_images(doc) -> list:
    """Check alt text on floating (anchored) images, shapes, and drawings.

    python-docx's inline_shapes only covers <wp:inline>. Floating images use
    <wp:anchor> and are invisible to that API. We walk the paragraph XML
    directly to find wp:anchor elements and check their docPr descr/title.
    Also catches SmartArt, grouped shapes, and canvases.
    """
    findings = []
    missing_alt = []
    poor_quality = []

    for para in doc.paragraphs:
        body = para._p
        # wp:anchor elements (floating images/shapes)
        for anchor in body.iter(f"{{{WP_NS}}}anchor"):
            doc_prs = anchor.findall(f".//{{{WP_NS}}}docPr")
            if not doc_prs:
                # Also check a:graphicData children for pic:pic docPr
                doc_prs = anchor.findall(f".//{{{A_NS}}}graphicData//{{{PIC_NS}}}cNvPr")
            if not doc_prs:
                doc_prs = anchor.findall(f".//{{{A_NS}}}graphicData//{{{WPS_NS}}}cNvSpPr/..")

            for doc_pr in doc_prs:
                descr = doc_pr.get("descr", "")
                title_attr = doc_pr.get("title", "")
                name = doc_pr.get("name", "floating shape")
                if not descr.strip() and not title_attr.strip():
                    missing_alt.append(name)
                elif descr.strip() and len(descr.strip()) < 4:
                    poor_quality.append(f"{name}: '{descr}'")

            if not doc_prs:
                # Anchor exists but no docPr found — flag it
                missing_alt.append("(floating element, no docPr)")

    if missing_alt:
        findings.append({
            "rule": "DOCX-IMG.ALT_FLOAT",
            "severity": "Error",
            "confidence": "High",
            "wcag": "1.1.1",
            "message": (
                f"{len(missing_alt)} floating image(s)/shape(s) have no alt text: "
                f"{missing_alt[:5]}{'...' if len(missing_alt) > 5 else ''}"
            ),
            "fix": (
                "Right-click the floating image/shape → View Alt Text → "
                "enter a description. If decorative, check 'Mark as decorative'. "
                "Note: floating images are commonly missed by accessibility checkers."
            ),
        })

    if poor_quality:
        findings.append({
            "rule": "DOCX-IMG.ALT_QUALITY",
            "severity": "Warning",
            "confidence": "Medium",
            "wcag": "1.1.1",
            "message": (
                f"{len(poor_quality)} floating image(s) have very short alt text: "
                f"{poor_quality[:5]}"
            ),
            "fix": "Review alt text to ensure it meaningfully describes the image content.",
        })

    return findings


def check_tables(doc) -> list:
    findings = []
    for t_idx, table in enumerate(doc.tables, 1):
        has_header = False
        has_merged = False
        try:
            # Check if first row has header style or repeat header row set
            if table.rows:
                first_row = table.rows[0]
                # Check for table header row repeat (trPr/tblHeader)
                tr_el = first_row._tr
                tbl_header = tr_el.find(f"{{{W_NS}}}trPr/{{{W_NS}}}tblHeader")
                if tbl_header is not None:
                    has_header = True
                # Check cell style names
                for cell in first_row.cells:
                    for para in cell.paragraphs:
                        style = para.style.name if para.style else ""
                        if "header" in style.lower() or "heading" in style.lower():
                            has_header = True

            # Detect merged cells (gridSpan > 1 or vMerge)
            for row in table.rows:
                for cell in row.cells:
                    tc_el = cell._tc
                    grid_span = tc_el.find(f"{{{W_NS}}}tcPr/{{{W_NS}}}gridSpan")
                    v_merge = tc_el.find(f"{{{W_NS}}}tcPr/{{{W_NS}}}vMerge")
                    if grid_span is not None or v_merge is not None:
                        has_merged = True

        except Exception:
            pass

        if not has_header:
            findings.append({
                "rule": "DOCX-TABLE.HEADERS",
                "severity": "Warning",
                "confidence": "Medium",
                "wcag": "1.3.1",
                "message": (
                    f"Table {t_idx}: No header row detected. "
                    "Without marking the header row, screen readers cannot identify column headers."
                ),
                "fix": (
                    "Click in the first row → Table Design tab → check 'Header Row'. "
                    "Also: Layout tab → Repeat Header Rows to repeat on multi-page tables."
                ),
            })

        if has_merged:
            findings.append({
                "rule": "DOCX-TABLE.MERGE",
                "severity": "Warning",
                "confidence": "High",
                "wcag": "1.3.1",
                "message": (
                    f"Table {t_idx}: Merged cells detected. "
                    "These create complex table structures that many screen readers struggle with."
                ),
                "fix": "Simplify the table to avoid merged cells where possible.",
            })

    return findings


def check_hyperlinks(doc) -> list:
    findings = []
    vague_links = []

    # Traverse XML for hyperlink elements
    for para in doc.paragraphs:
        for hyper in para._p.findall(f".//{{{W_NS}}}hyperlink"):
            text_nodes = hyper.findall(f".//{{{W_NS}}}t")
            link_text = "".join(t.text or "" for t in text_nodes).strip()
            if link_text and VAGUE_LINK_PATTERNS.match(link_text):
                vague_links.append(link_text)

    if vague_links:
        findings.append({
            "rule": "DOCX-LINK.DESCRIPTIVE",
            "severity": "Warning",
            "confidence": "High",
            "wcag": "2.4.4",
            "message": (
                f"{len(vague_links)} hyperlink(s) have non-descriptive text: "
                f"{list(set(vague_links))[:5]}"
            ),
            "fix": (
                "Replace generic link text ('click here', 'read more') with "
                "descriptive text that makes sense out of context, e.g. "
                "'Download the Medication Authorization Form (PDF)'."
            ),
        })

    return findings


def check_empty_paragraphs(doc) -> list:
    findings = []
    empty_count = 0
    for para in doc.paragraphs:
        if not para.text.strip() and para.style.name == "Normal":
            empty_count += 1
    if empty_count > 3:
        findings.append({
            "rule": "DOCX-STRUCT.EMPTYPARA",
            "severity": "Info",
            "confidence": "High",
            "message": (
                f"{empty_count} empty paragraphs detected. "
                "These are often used for visual spacing and can cause screen readers "
                "to announce 'blank' repeatedly."
            ),
            "fix": (
                "Use paragraph spacing (Format → Paragraph → Spacing Before/After) "
                "instead of empty paragraphs for visual spacing."
            ),
        })
    return findings


# ---------------------------------------------------------------------------
# List semantics check
# ---------------------------------------------------------------------------

_MANUAL_BULLET_RE = re.compile(
    r"^[\s]*[•●○■▪▸►\-–—\*]\s",
)
_MANUAL_NUMBER_RE = re.compile(
    r"^[\s]*(?:\d{1,3}[.)]\s|[a-zA-Z][.)]\s|[ivxIVX]+[.)]\s)",
)


def check_list_semantics(doc) -> list:
    """Detect paragraphs that look like manual lists but don't use Word list styles.

    Manual bullets/numbering (typed characters like '- item' or '1. item') are
    invisible to screen readers as list items. Word's built-in list styles
    (List Bullet, List Number, or any style with numPr) produce proper list
    semantics that assistive technology can navigate.
    """
    findings = []
    manual_bullet_count = 0
    manual_number_count = 0

    for para in doc.paragraphs:
        text = para.text
        if not text.strip():
            continue
        style_name = para.style.name if para.style else ""

        # Check if already using a proper list style
        is_list_style = any(kw in style_name.lower() for kw in
                           ("list", "bullet", "number", "toc"))
        # Also check numPr in XML (covers custom numbered styles)
        p_elem = para._p
        num_pr = p_elem.find(f".//{{{W_NS}}}numPr")
        if is_list_style or num_pr is not None:
            continue  # properly styled

        if _MANUAL_BULLET_RE.match(text):
            manual_bullet_count += 1
        elif _MANUAL_NUMBER_RE.match(text):
            manual_number_count += 1

    total = manual_bullet_count + manual_number_count
    if total >= 3:
        details = []
        if manual_bullet_count:
            details.append(f"{manual_bullet_count} bullet-like")
        if manual_number_count:
            details.append(f"{manual_number_count} numbered-like")
        findings.append({
            "rule": "DOCX-LIST.SEMANTIC",
            "severity": "Warning",
            "confidence": "Medium",
            "wcag": "1.3.1",
            "message": (
                f"{total} paragraph(s) appear to be manual lists ({', '.join(details)}) "
                "without using Word list styles. Screen readers cannot identify or "
                "navigate these as lists."
            ),
            "fix": (
                "Select the list items → Home tab → use the Bullets or "
                "Numbering button (or apply 'List Bullet' / 'List Number' styles). "
                "This creates proper list semantics for assistive technology."
            ),
        })

    return findings


# ---------------------------------------------------------------------------
# Nested tables check
# ---------------------------------------------------------------------------

def check_nested_tables(doc) -> list:
    """Detect tables nested inside other tables.

    Nested tables are extremely difficult for screen readers to navigate.
    The cell-by-cell navigation model breaks down when a cell contains
    another table, often causing users to lose their place entirely.
    """
    findings = []
    nested_count = 0

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                # Check if cell contains a nested table via XML
                tc_el = cell._tc
                nested_tbls = tc_el.findall(f".//{{{W_NS}}}tbl")
                if nested_tbls:
                    nested_count += len(nested_tbls)

    if nested_count:
        findings.append({
            "rule": "DOCX-TABLE.NESTED",
            "severity": "Error",
            "confidence": "High",
            "wcag": "1.3.1",
            "message": (
                f"{nested_count} nested table(s) found. "
                "Nested tables are extremely difficult for screen readers to navigate. "
                "Users lose track of which table and cell they are in."
            ),
            "fix": (
                "Restructure the content to avoid nested tables. "
                "Use heading levels, separate tables, or paragraph text instead."
            ),
        })

    return findings


# ---------------------------------------------------------------------------
# Table of Contents check (long documents)
# ---------------------------------------------------------------------------

def check_toc(doc) -> list:
    """Check for a Table of Contents in long documents.

    Documents with many headings benefit from a generated TOC for navigation.
    The TOC must be a real Word TOC field (w:sdt with TOC), not manual text.
    """
    findings = []
    heading_count = sum(
        1 for p in doc.paragraphs
        if p.style and p.style.name in HEADING_STYLES
    )
    if heading_count < 5:
        return findings  # short doc, TOC not expected

    # Look for TOC structured document tag or TOC field code
    body_xml = doc.element.body
    # SDT-based TOC (modern Word)
    has_toc = False
    for sdt in body_xml.iter(f"{{{W_NS}}}sdt"):
        sdt_pr = sdt.find(f"{{{W_NS}}}sdtPr")
        if sdt_pr is not None:
            doc_part = sdt_pr.find(f"{{{W_NS}}}docPartObj")
            if doc_part is not None:
                gallery = doc_part.find(f"{{{W_NS}}}docPartGallery")
                if gallery is not None and "toc" in (gallery.get(f"{{{W_NS}}}val") or "").lower():
                    has_toc = True
                    break
    # Also check for field-code-based TOC
    if not has_toc:
        for fld in body_xml.iter(f"{{{W_NS}}}fldChar"):
            fld_type = fld.get(f"{{{W_NS}}}fldCharType")
            if fld_type == "begin":
                # Look at next sibling instrText
                parent = fld.getparent()
                for sibling in parent.itersiblings():
                    instr = sibling.find(f"{{{W_NS}}}instrText")
                    if instr is not None and instr.text and "TOC" in instr.text.upper():
                        has_toc = True
                        break
                if has_toc:
                    break

    if not has_toc:
        findings.append({
            "rule": "DOCX-STRUCT.TOC",
            "severity": "Warning",
            "confidence": "Medium",
            "wcag": "2.4.5",
            "message": (
                f"Document has {heading_count} headings but no Table of Contents. "
                "Long documents benefit from a generated TOC for navigation."
            ),
            "fix": (
                "Place cursor where you want the TOC → References tab → "
                "Table of Contents → choose an Automatic style. "
                "This creates a navigable TOC linked to your heading styles."
            ),
        })

    return findings


# ---------------------------------------------------------------------------
# Tracked changes detection
# ---------------------------------------------------------------------------

def check_tracked_changes(doc) -> list:
    """Detect unresolved tracked changes (insertions, deletions, moves).

    Tracked changes cause screen readers to announce revision metadata
    (author, date, change type) for every change, making the document
    extremely noisy and difficult to navigate.
    """
    findings = []
    body_xml = doc.element.body
    insertions = len(body_xml.findall(f".//{{{W_NS}}}ins"))
    deletions = len(body_xml.findall(f".//{{{W_NS}}}del"))
    moves = len(body_xml.findall(f".//{{{W_NS}}}moveTo"))
    total = insertions + deletions + moves

    if total > 0:
        details = []
        if insertions:
            details.append(f"{insertions} insertion(s)")
        if deletions:
            details.append(f"{deletions} deletion(s)")
        if moves:
            details.append(f"{moves} move(s)")
        findings.append({
            "rule": "DOCX-REVIEW.TRACKED",
            "severity": "Warning",
            "confidence": "High",
            "wcag": "1.3.1",
            "message": (
                f"Document contains {total} unresolved tracked change(s): "
                f"{', '.join(details)}. Screen readers announce revision metadata "
                "for every change, making the document extremely difficult to read."
            ),
            "fix": (
                "Accept or reject all tracked changes before publishing: "
                "Review tab → Accept → Accept All Changes in Document."
            ),
        })

    return findings


# ---------------------------------------------------------------------------
# Header/footer content check
# ---------------------------------------------------------------------------

def check_headers_footers(doc) -> list:
    """Check for meaningful info-bearing content in headers/footers.

    Content in headers/footers (other than page numbers) may not be read by
    screen readers in all contexts, and users may miss important information
    placed only in these areas.
    """
    findings = []
    info_bearing = []

    for section in doc.sections:
        for hf_attr, hf_label in [
            ("header", "header"), ("footer", "footer"),
            ("first_page_header", "first-page header"),
            ("first_page_footer", "first-page footer"),
            ("even_page_header", "even-page header"),
            ("even_page_footer", "even-page footer"),
        ]:
            try:
                hf = getattr(section, hf_attr)
                if not hf.is_linked_to_previous:
                    text = " ".join(p.text.strip() for p in hf.paragraphs).strip()
                    # Ignore empty or just page-number-like content
                    if text and not re.match(r"^[\d\s/\-–—pageofPage]*$", text, re.IGNORECASE):
                        if len(text) > 20:
                            info_bearing.append(f"{hf_label}: '{text[:40]}...'")
            except Exception:
                pass

    if info_bearing:
        findings.append({
            "rule": "DOCX-STRUCT.HDRFTR",
            "severity": "Info",
            "confidence": "Medium",
            "wcag": "1.3.2",
            "message": (
                f"{len(info_bearing)} header/footer section(s) contain substantial text "
                "that may not be consistently announced by screen readers. "
                f"Examples: {info_bearing[:3]}"
            ),
            "fix": (
                "Ensure any important information in headers/footers is also "
                "present in the main document body. Header/footer content "
                "may be skipped by assistive technology depending on the reading mode."
            ),
        })

    return findings


# ---------------------------------------------------------------------------
# Footnotes / endnotes check
# ---------------------------------------------------------------------------

def check_footnotes(docx_path: Path) -> list:
    """Check footnote/endnote count and flag very long notes.

    Footnotes and endnotes are generally accessible in Word, but documents
    with excessive notes can be difficult to navigate. Very long footnotes
    may break reading flow significantly.
    """
    findings = []
    footnote_count = 0
    endnote_count = 0
    long_notes = 0

    try:
        with ZipFile(str(docx_path), "r") as zf:
            # Count footnotes
            if "word/footnotes.xml" in zf.namelist():
                tree = etree.parse(zf.open("word/footnotes.xml"))
                root = tree.getroot()
                # Skip separator/continuation footnotes (type="separator" or "continuationSeparator")
                for fn in root.findall(f".//{{{W_NS}}}footnote"):
                    fn_type = fn.get(f"{{{W_NS}}}type", "")
                    if fn_type in ("separator", "continuationSeparator"):
                        continue
                    footnote_count += 1
                    # Check note length
                    text = "".join(fn.itertext()).strip()
                    if len(text) > 300:
                        long_notes += 1

            # Count endnotes
            if "word/endnotes.xml" in zf.namelist():
                tree = etree.parse(zf.open("word/endnotes.xml"))
                root = tree.getroot()
                for en in root.findall(f".//{{{W_NS}}}endnote"):
                    en_type = en.get(f"{{{W_NS}}}type", "")
                    if en_type in ("separator", "continuationSeparator"):
                        continue
                    endnote_count += 1
                    text = "".join(en.itertext()).strip()
                    if len(text) > 300:
                        long_notes += 1

    except (BadZipFile, Exception):
        pass

    total = footnote_count + endnote_count
    if total > 20:
        findings.append({
            "rule": "DOCX-STRUCT.FOOTNOTES",
            "severity": "Info",
            "confidence": "Medium",
            "wcag": "2.4.1",
            "message": (
                f"Document has {footnote_count} footnote(s) and {endnote_count} endnote(s) "
                f"({total} total). Heavy use of notes can interrupt reading flow "
                "for screen reader users who navigate to each note reference."
            ),
            "fix": (
                "Consider consolidating notes or converting to inline references "
                "where possible. Ensure all note content is essential."
            ),
        })

    if long_notes:
        findings.append({
            "rule": "DOCX-STRUCT.FOOTNOTES",
            "severity": "Warning",
            "confidence": "Medium",
            "wcag": "2.4.1",
            "message": (
                f"{long_notes} footnote/endnote(s) are unusually long (>300 chars). "
                "Long notes significantly interrupt the reading flow when a screen "
                "reader navigates to the note and back."
            ),
            "fix": (
                "Move lengthy notes into the main document body, an appendix, "
                "or link to supplementary material."
            ),
        })

    return findings


# ---------------------------------------------------------------------------
# Main scan function
# ---------------------------------------------------------------------------

def scan_word(docx_path: Path) -> dict:
    result = {
        "file": docx_path.name,
        "path": str(docx_path),
        "format": "docx",
        "findings": [],
        "errors": [],
    }

    try:
        doc = Document(str(docx_path))
        result["findings"] += check_metadata(doc)
        result["findings"] += check_headings(doc)
        result["findings"] += check_images(doc)
        result["findings"] += check_floating_images(doc)
        result["findings"] += check_tables(doc)
        result["findings"] += check_nested_tables(doc)
        result["findings"] += check_hyperlinks(doc)
        result["findings"] += check_list_semantics(doc)
        result["findings"] += check_empty_paragraphs(doc)
        result["findings"] += check_toc(doc)
        result["findings"] += check_tracked_changes(doc)
        result["findings"] += check_headers_footers(doc)
        result["findings"] += check_footnotes(docx_path)
    except Exception as e:
        result["errors"].append(f"Could not open document: {e}")

    return result


def scan_folder(folder: Path) -> list:
    return [scan_word(p) for p in sorted(folder.glob("*.docx"))]


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Scan Word (.docx) documents for accessibility issues."
    )
    parser.add_argument("path", help="DOCX file or folder")
    parser.add_argument("--output", help="Write JSON output here")
    parser.add_argument("--json", action="store_true", help="Print JSON to stdout")
    args = parser.parse_args()

    target = Path(args.path)
    if target.is_dir():
        results = scan_folder(target)
    elif target.is_file():
        results = [scan_word(target)]
    else:
        sys.exit(f"ERROR: Path not found: {target}")

    if args.output:
        Path(args.output).write_text(json.dumps(results, indent=2), encoding="utf-8")
        print(f"JSON written to {args.output}")
    elif args.json:
        print(json.dumps(results, indent=2))
    else:
        for r in results:
            errs = sum(1 for f in r["findings"] if f["severity"] == "Error")
            warns = sum(1 for f in r["findings"] if f["severity"] == "Warning")
            print(f"\n{r['file']}: {errs} errors, {warns} warnings")
            for f in r["findings"]:
                icon = "❌" if f["severity"] == "Error" else "⚠️" if f["severity"] == "Warning" else "ℹ️"
                print(f"  {icon} [{f['rule']}] {f['message'][:100]}")


if __name__ == "__main__":
    sys.exit(main())
