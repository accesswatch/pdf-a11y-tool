"""
fix_word.py — Word Document Accessibility Auto-Fixer
=====================================================
Applies deterministic accessibility fixes to .docx files:
  - Missing document title
  - Missing document language
  - Missing author metadata
  - Missing table header rows
  - Missing alt text placeholders on images

Always creates a backup and writes to a new -fixed.docx file.
Never overwrites the original.

Usage (standalone):
    python tools/fix_word.py input.docx [--output fixed.docx] [--rules DOCX-META.TITLE,DOCX-META.LANG]

This file is PERMANENT — do not delete.
"""

import copy
import json
import shutil
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# ---------------------------------------------------------------------------
# Fix functions
# ---------------------------------------------------------------------------

def _fix_title(doc: "Document", title: str = "") -> dict:
    """Set document title if missing."""
    current = doc.core_properties.title or ""
    if current.strip():
        return {"rule": "DOCX-META.TITLE", "status": "skipped", "reason": "Title already set"}
    doc.core_properties.title = title or "TODO: Add descriptive document title"
    return {"rule": "DOCX-META.TITLE", "status": "fixed", "detail": f"Set title to '{doc.core_properties.title}'"}


def _fix_author(doc: "Document", author: str = "") -> dict:
    """Set document author if missing."""
    current = doc.core_properties.author or ""
    if current.strip():
        return {"rule": "DOCX-META.AUTHOR", "status": "skipped", "reason": "Author already set"}
    doc.core_properties.author = author or "TODO: Add author name"
    return {"rule": "DOCX-META.AUTHOR", "status": "fixed", "detail": f"Set author to '{doc.core_properties.author}'"}


def _fix_language(doc: "Document", lang: str = "en-US") -> dict:
    """Set document language via styles.xml if not already set."""
    try:
        styles_elem = doc.styles.element
        rpr_default = styles_elem.find(qn("w:docDefaults"))
        if rpr_default is not None:
            rpr = rpr_default.find(f".//{qn('w:rPr')}")
            if rpr is not None:
                existing_lang = rpr.find(qn("w:lang"))
                if existing_lang is not None:
                    val = existing_lang.get(qn("w:val"), "")
                    if val:
                        return {"rule": "DOCX-META.LANG", "status": "skipped", "reason": f"Language already set: {val}"}
                # Add or update lang element
                if existing_lang is None:
                    existing_lang = rpr.makeelement(qn("w:lang"), {})
                    rpr.append(existing_lang)
                existing_lang.set(qn("w:val"), lang)
                existing_lang.set(qn("w:eastAsia"), lang)
                existing_lang.set(qn("w:bidi"), lang)
                return {"rule": "DOCX-META.LANG", "status": "fixed", "detail": f"Set language to {lang}"}
        return {"rule": "DOCX-META.LANG", "status": "failed", "reason": "Could not locate docDefaults in styles.xml"}
    except Exception as e:
        return {"rule": "DOCX-META.LANG", "status": "failed", "reason": str(e)}


def _fix_table_headers(doc: "Document") -> list[dict]:
    """Set header row property on all tables missing it."""
    results = []
    for i, table in enumerate(doc.tables, 1):
        first_row = table.rows[0]
        tr = first_row._tr
        trPr = tr.get_or_add_trPr()
        existing = trPr.find(qn("w:tblHeader"))
        if existing is not None:
            continue
        header_elem = copy.deepcopy(tr.makeelement(qn("w:tblHeader"), {}))
        header_elem.set(qn("w:val"), "true")
        trPr.append(header_elem)
        results.append({
            "rule": "DOCX-TABLE.HEADERS",
            "status": "fixed",
            "detail": f"Set header row on table {i}",
        })
    if not results:
        results.append({"rule": "DOCX-TABLE.HEADERS", "status": "skipped", "reason": "All tables already have headers or no tables found"})
    return results


def _fix_image_alt_text(doc: "Document", use_alt_text_generator: bool = False,
                        alt_text_models: list[str] | None = None,
                        alt_text_context: str = "") -> list[dict]:
    """Add alt text to images missing it.

    When use_alt_text_generator is True, generates real alt text via
    the vision LLM. Otherwise, adds TODO placeholders.
    """
    results = []
    body = doc.element.body
    drawings = body.findall(f".//{qn('w:drawing')}")
    img_count = 0
    for drawing in drawings:
        for docPr in drawing.findall(f".//{qn('wp:docPr')}"):
            descr = docPr.get("descr", "").strip()
            name = docPr.get("name", "Image")
            if not descr:
                img_count += 1

                if use_alt_text_generator:
                    alt = _generate_alt_for_docx_image(
                        doc, drawing, name, alt_text_models, alt_text_context
                    )
                    if alt:
                        docPr.set("descr", alt)
                        results.append({
                            "rule": "DOCX-IMG.ALT",
                            "status": "fixed",
                            "detail": f"Generated alt text for '{name}': '{alt[:80]}...'",
                        })
                        continue

                # Fallback: placeholder
                placeholder = f"TODO: Describe {name}"
                docPr.set("descr", placeholder)
                results.append({
                    "rule": "DOCX-IMG.ALT",
                    "status": "fixed",
                    "detail": f"Added placeholder alt text to '{name}': '{placeholder}'",
                })
    if not results:
        results.append({"rule": "DOCX-IMG.ALT", "status": "skipped", "reason": "All images have alt text or no images found"})
    return results


def _generate_alt_for_docx_image(
    doc, drawing, name: str,
    models: list[str] | None, context: str,
) -> str | None:
    """Try to generate alt text using the alt-text generator."""
    try:
        from alt_text.client import generate_for_image
    except ImportError:
        return None

    try:
        # Extract image bytes from the drawing's blip reference
        from docx.oxml.ns import qn as _qn
        blip = drawing.find(f".//{_qn('a:blip')}")
        if blip is None:
            return None

        r_embed = blip.get(f"{{{_NS_R}}}embed", "")
        if not r_embed:
            return None

        # Get the image part from the document relationships
        part = doc.part
        rel = part.rels.get(r_embed)
        if rel is None:
            return None

        image_bytes = rel.target_part.blob
        content_type = rel.target_part.content_type or "image/png"

        result = generate_for_image(
            image_bytes=image_bytes,
            mime_type=content_type,
            model=(models or ["openai/gpt-4.1"])[0],
            profile="auto",
            source_format="docx",
            image_name=name,
            context=context,
        )
        return result.concise_alt if result.concise_alt else None
    except Exception:
        return None


_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# ---------------------------------------------------------------------------
# Main fix orchestrator
# ---------------------------------------------------------------------------

_ALL_RULES = {
    "DOCX-META.TITLE": _fix_title,
    "DOCX-META.LANG": _fix_language,
    "DOCX-META.AUTHOR": _fix_author,
    "DOCX-TABLE.HEADERS": _fix_table_headers,
    "DOCX-IMG.ALT": _fix_image_alt_text,
}


def fix_word(file_path: Path, output_path: Path | None = None,
             rules: list[str] | None = None,
             title: str = "", author: str = "", lang: str = "en-US") -> dict:
    """Apply accessibility fixes to a Word document.

    Args:
        file_path: Path to the original .docx file.
        output_path: Where to save the fixed file. Defaults to <name>-fixed.docx.
        rules: Optional list of rule IDs to fix. None = fix all.
        title: Custom title text. Empty = placeholder.
        author: Custom author text. Empty = placeholder.
        lang: BCP 47 language tag. Default: en-US.

    Returns:
        Dict with keys: file, output, backup, fixes (list), summary.
    """
    if not HAS_DOCX:
        return {"error": "python-docx is not installed. Run: pip install python-docx"}

    file_path = Path(file_path)
    if not file_path.exists():
        return {"error": f"File not found: {file_path}"}

    if output_path is None:
        output_path = file_path.with_stem(file_path.stem + "-fixed")
    else:
        output_path = Path(output_path)

    # Create backup
    backup_path = file_path.with_stem(file_path.stem + "-backup")
    shutil.copy2(file_path, backup_path)

    doc = Document(str(file_path))
    all_fixes = []
    target_rules = rules if rules else list(_ALL_RULES.keys())

    for rule_id in target_rules:
        if rule_id not in _ALL_RULES:
            all_fixes.append({"rule": rule_id, "status": "unknown", "reason": f"Rule {rule_id} not supported for auto-fix"})
            continue

        func = _ALL_RULES[rule_id]
        if rule_id == "DOCX-META.TITLE":
            result = func(doc, title)
        elif rule_id == "DOCX-META.AUTHOR":
            result = func(doc, author)
        elif rule_id == "DOCX-META.LANG":
            result = func(doc, lang)
        else:
            result = func(doc)

        if isinstance(result, list):
            all_fixes.extend(result)
        else:
            all_fixes.append(result)

    doc.save(str(output_path))

    fixed_count = sum(1 for f in all_fixes if f["status"] == "fixed")
    skipped_count = sum(1 for f in all_fixes if f["status"] == "skipped")
    failed_count = sum(1 for f in all_fixes if f["status"] == "failed")

    return {
        "file": str(file_path),
        "output": str(output_path),
        "backup": str(backup_path),
        "fixes": all_fixes,
        "summary": {
            "fixed": fixed_count,
            "skipped": skipped_count,
            "failed": failed_count,
            "total_rules_attempted": len(target_rules),
        },
    }


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Fix Word document accessibility issues")
    parser.add_argument("input", help="Path to .docx file")
    parser.add_argument("--output", help="Output file path")
    parser.add_argument("--rules", help="Comma-separated rule IDs to fix")
    parser.add_argument("--title", default="", help="Document title")
    parser.add_argument("--author", default="", help="Document author")
    parser.add_argument("--lang", default="en-US", help="Document language (BCP 47)")
    args = parser.parse_args()

    rule_list = args.rules.split(",") if args.rules else None
    out = args.output if args.output else None
    result = fix_word(Path(args.input), Path(out) if out else None, rule_list, args.title, args.author, args.lang)
    print(json.dumps(result, indent=2))
