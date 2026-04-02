"""
extractor_pptx.py — PowerPoint Image Extraction
=================================================
Extracts images from .pptx files using python-pptx.

This file is PERMANENT — do not delete.
"""

from __future__ import annotations

import hashlib
import logging
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from .config import AltTextConfig

from .client import ExtractedImage

log = logging.getLogger(__name__)

_MIME_MAP = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".bmp": "image/bmp",
    ".tiff": "image/tiff",
    ".emf": "image/x-emf",
    ".wmf": "image/x-wmf",
    ".svg": "image/svg+xml",
}


def extract_images(path: Path, config: "AltTextConfig") -> list[ExtractedImage]:
    """Extract images from a PowerPoint presentation.

    Args:
        path: Path to the .pptx file.
        config: AltTextConfig controlling extraction behavior.

    Returns:
        List of ExtractedImage objects.
    """
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        log.warning("python-pptx not installed — cannot extract PPTX images")
        return []

    images: list[ExtractedImage] = []
    seen_hashes: set[str] = set()

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        log.error("Failed to open pptx %s: %s", path, exc)
        return []

    img_idx = 0
    for slide_num, slide in enumerate(prs.slides, 1):
        # Gather slide text for context
        slide_text = _get_slide_text(slide)

        for shape in slide.shapes:
            if not hasattr(shape, "image"):
                continue

            try:
                image = shape.image
            except Exception:
                continue

            name = shape.name or f"Slide{slide_num}_Image{img_idx}"

            # Get existing alt text
            existing_alt = _get_shape_alt(shape)
            if config.skip_existing_alt and existing_alt:
                continue

            try:
                image_bytes = image.blob
            except Exception:
                continue

            if not image_bytes:
                continue

            # Skip duplicates
            if config.skip_duplicate_images:
                img_hash = hashlib.md5(image_bytes).hexdigest()
                if img_hash in seen_hashes:
                    continue
                seen_hashes.add(img_hash)

            # MIME type
            ext = "." + (image.content_type.split("/")[-1] if image.content_type else "png")
            mime_type = _MIME_MAP.get(ext, image.content_type or "image/png")

            images.append(ExtractedImage(
                image_bytes=image_bytes,
                mime_type=mime_type,
                page_number=slide_num,
                image_index=img_idx,
                image_name=name,
                existing_alt=existing_alt,
                surrounding_text=slide_text,
                slide_number=slide_num,
                shape_name=name,
            ))
            img_idx += 1

    return images


def _get_shape_alt(shape) -> str | None:
    """Get alt text from a pptx shape's cNvPr descr attribute."""
    try:
        _NS_PML = "http://schemas.openxmlformats.org/presentationml/2006/main"
        _NS_DML = "http://schemas.openxmlformats.org/drawingml/2006/main"

        # Try nvPicPr/cNvPr
        for ns in (_NS_PML, _NS_DML):
            nvPicPr = shape._element.find(f".//{{{ns}}}nvPicPr")
            if nvPicPr is not None:
                cNvPr = nvPicPr.find(f"{{{ns}}}cNvPr")
                if cNvPr is not None:
                    descr = cNvPr.get("descr", "").strip()
                    if descr:
                        return descr

        # Direct element check
        for attr_name in ("descr", "title"):
            val = shape._element.get(attr_name, "").strip()
            if val:
                return val
    except Exception:
        pass
    return None


def _get_slide_text(slide) -> str:
    """Get all text content from a slide, truncated."""
    parts: list[str] = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    parts.append(text)
    return " ".join(parts)[:500]
